#!/usr/bin/env python3
"""
Bulk-ingest the full all_trainees_view.xlsx export into at_rows via the deployed
Worker's token-gated /api/mis-sync/ingest route.

Why this over the gateway drain: the MIS gateway is slow/unstable, but the client
exported the complete sheet (~800k rows). Ingesting the file closes the 32k gap
in one reliable pass; the gateway sync then only needs the LATEST new rows
(cheap freshness pass), so the dashboard always picks up the newest data without
depending on the flaky deep-page backfill.

CRITICAL — dedup_key must match gateway-ingested rows:
  The at_rows dedup key hashes (participant_name, participant_id, group_id,
  training_type, activity_date, data_collector, group_name, sex, district,
  subcounty, Parish, Village, Disability_status, Employment_status,
  Employment_sector, Do_for_living) each .trim()'d. The gateway sends
  activity_date as ISO "2026-06-03T00:00:00.000" which misRowToRecord slices to
  "2026-06-03". The Excel shows "Wednesday, 3 June 2026", so we convert it to the
  SAME "YYYY-MM-DD" here; every other field is passed as a string (None -> "").
  Result: rows already present dedup to no-ops; only the missing ones insert.

Usage:
  python3 scripts/ingest_xlsx.py [--start N] [--limit M] [--batch B] [--dry]
Env: DEPLOY, TOKEN (defaults match the deployed project).
"""
import argparse, datetime, json, os, sys, time, urllib.request

XLSX = os.environ.get("XLSX", "/home/user/uploaded_files/all_trainees_view (2).xlsx")
DEPLOY = os.environ.get("DEPLOY", "https://shg-data-cleaner.pages.dev")
TOKEN = os.environ.get("TOKEN", "shg-fix-2026")

HEADER = ['_id','participant_name','participant_id','group_id','training_type',
          'activity_date','data_collector','group_name','sex','district',
          'subcounty','Parish','Village','Disability_status','Employment_status',
          'Employment_sector','Do_for_living']

# Excel serves dates as either a python datetime (if cell is date-typed) or a
# human string like "Wednesday, 3 June 2026". Normalize BOTH to YYYY-MM-DD.
def to_iso_date(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    if not s:
        return ""
    # already ISO-ish?
    if len(s) >= 10 and s[4] == '-' and s[7] == '-':
        return s[:10]
    # try "Wednesday, 3 June 2026" and a few common variants
    for fmt in ("%A, %d %B %Y", "%a, %d %B %Y", "%d %B %Y", "%B %d, %Y",
                "%d/%m/%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    # unknown format: leave verbatim (misRowToRecord will keep it as-is; a
    # mismatch just means a potential dup, never data loss)
    return s

def cell(v):
    if v is None:
        return ""
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%Y-%m-%d")
    return str(v)

def post_batch(rows):
    body = json.dumps({"rows": rows}).encode("utf-8")
    req = urllib.request.Request(
        f"{DEPLOY}/api/mis-sync/ingest?token={TOKEN}",
        data=body,
        # A browser-like UA is REQUIRED: Cloudflare's edge blocks the default
        # "Python-urllib/*" agent with a 403 before the request reaches our app.
        headers={"Content-Type": "application/json",
                 "User-Agent": "Mozilla/5.0 (xlsx-ingest)"},
        method="POST")
    for attempt in range(1, 6):
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                return json.loads(r.read().decode("utf-8"))
        except Exception as e:
            if attempt == 5:
                return {"error": str(e)[:160]}
            time.sleep(attempt * 3)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--start", type=int, default=0, help="skip this many data rows")
    ap.add_argument("--limit", type=int, default=0, help="max data rows (0=all)")
    ap.add_argument("--batch", type=int, default=500)
    ap.add_argument("--dry", action="store_true", help="parse only, no POST")
    args = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = list(next(it))
    if header[:len(HEADER)] != HEADER:
        print("WARN: header mismatch\n got:", header, file=sys.stderr)

    seen = 0            # data rows scanned
    sent = 0            # rows POSTed
    up_total = 0        # upserted (new) rows
    batch = []
    t0 = time.time()

    def flush():
        nonlocal sent, up_total, batch
        if not batch:
            return
        if args.dry:
            up_total += 0
        else:
            res = post_batch(batch)
            if "error" in res:
                print(f"  batch@{sent} INGEST ERROR: {res['error']}", file=sys.stderr)
            else:
                up = int(res.get("upserted", 0))
                up_total += up
                at = res.get("atRowsCount", "?")
                print(f"  rows {sent}-{sent+len(batch)-1}: upserted={up} at_rows={at} cum_new={up_total} ({round(time.time()-t0)}s)")
        sent += len(batch)
        batch = []

    for row in it:
        if seen < args.start:
            seen += 1
            continue
        if args.limit and (seen - args.start) >= args.limit:
            break
        seen += 1
        d = {}
        for i, name in enumerate(HEADER):
            v = row[i] if i < len(row) else None
            d[name] = to_iso_date(v) if name == "activity_date" else cell(v)
        batch.append(d)
        if len(batch) >= args.batch:
            flush()

    flush()
    wb.close()
    print(f"DONE scanned={seen-args.start} sent={sent} new_upserted={up_total} elapsed={round(time.time()-t0)}s")

if __name__ == "__main__":
    main()
